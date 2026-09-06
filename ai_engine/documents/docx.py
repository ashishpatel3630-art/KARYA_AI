from pathlib import Path

from docx import Document


class DOCXParser:
    """Extract text from Word DOCX documents."""

    def parse(self, file_path: str) -> str:
        """
        Extract text from paragraphs and tables in a DOCX file.

        Args:
            file_path: Path to the DOCX file.

        Returns:
            Extracted document text.
        """

        path = Path(file_path)

        if not path.exists():
            raise FileNotFoundError(
                f"DOCX file not found: {file_path}"
            )

        if path.suffix.lower() != ".docx":
            raise ValueError(
                f"Expected a DOCX file, got: {path.suffix}"
            )

        document = Document(str(path))

        content = []

        # Extract normal paragraphs
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()

            if text:
                content.append(text)

        # Extract tables
        for table in document.tables:
            for row in table.rows:
                cells = []

                for cell in row.cells:
                    text = cell.text.strip()

                    if text:
                        cells.append(text)

                if cells:
                    content.append(" | ".join(cells))

        return "\n".join(content)
