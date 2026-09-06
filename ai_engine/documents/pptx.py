from pathlib import Path

from pptx import Presentation


class PPTXParser:
    """Extract text from PowerPoint PPTX presentations."""

    def parse(self, file_path: str) -> str:
        """
        Extract text from all slides in a PPTX file.

        Args:
            file_path: Path to the PPTX file.

        Returns:
            Presentation text as a single string.
        """

        path = Path(file_path)

        if not path.exists():
            raise FileNotFoundError(
                f"PPTX file not found: {file_path}"
            )

        if path.suffix.lower() != ".pptx":
            raise ValueError(
                f"Expected a PPTX file, got: {path.suffix}"
            )

        presentation = Presentation(str(path))

        content = []

        for slide_number, slide in enumerate(
            presentation.slides,
            start=1,
        ):
            content.append(f"[Slide: {slide_number}]")

            for shape in slide.shapes:
                if not hasattr(shape, "text"):
                    continue

                text = shape.text.strip()

                if text:
                    content.append(text)

        return "\n".join(content)
