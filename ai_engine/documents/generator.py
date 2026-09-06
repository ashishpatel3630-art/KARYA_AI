from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pptx import Presentation


class DocumentGenerator:
    """Generate business documents for KARYA AI."""

    def generate_docx(
        self,
        output_path: str,
        title: str,
        content: str,
    ) -> str:
        """
        Generate a DOCX document.

        Args:
            output_path: Destination path.
            title: Document title.
            content: Document body.

        Returns:
            Path to the generated document.
        """

        path = Path(output_path)

        path.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        document = Document()

        document.add_heading(
            title,
            level=1,
        )

        for paragraph in content.split("\n"):
            paragraph = paragraph.strip()

            if paragraph:
                document.add_paragraph(paragraph)

        document.save(str(path))

        return str(path)

    def generate_xlsx(
        self,
        output_path: str,
        sheet_name: str,
        headers: list[str],
        rows: list[list],
    ) -> str:
        """
        Generate an XLSX spreadsheet.

        Args:
            output_path: Destination path.
            sheet_name: Worksheet name.
            headers: Column headers.
            rows: Spreadsheet rows.

        Returns:
            Path to the generated spreadsheet.
        """

        path = Path(output_path)

        path.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        workbook = Workbook()

        worksheet = workbook.active
        worksheet.title = sheet_name

        worksheet.append(headers)

        for row in rows:
            worksheet.append(row)

        workbook.save(str(path))

        return str(path)

    def generate_pptx(
        self,
        output_path: str,
        title: str,
        slides: list[dict],
    ) -> str:
        """
        Generate a PPTX presentation.

        Args:
            output_path: Destination path.
            title: Presentation title.
            slides: List of slides containing title and content.

        Returns:
            Path to the generated presentation.
        """

        path = Path(output_path)

        path.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        presentation = Presentation()

        # Title slide
        title_slide = presentation.slides.add_slide(
            presentation.slide_layouts[0]
        )

        title_slide.shapes.title.text = title

        # Content slides
        for slide_data in slides:
            slide = presentation.slides.add_slide(
                presentation.slide_layouts[1]
            )

            slide.shapes.title.text = slide_data["title"]

            slide.placeholders[1].text = slide_data["content"]

        presentation.save(str(path))

        return str(path)