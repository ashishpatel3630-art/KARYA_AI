from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from documents.pptx import PPTXParser


def test_pptx_parser_reads_slides(tmp_path: Path):
	presentation_path = tmp_path / "test_document.pptx"
	presentation = Presentation()
	slide = presentation.slides.add_slide(
		presentation.slide_layouts[5]
	)
	textbox = slide.shapes.add_textbox(
		Inches(1),
		Inches(1),
		Inches(4),
		Inches(1),
	)
	textbox.text = "KARYA PPTX test"
	presentation.save(presentation_path)

	slides = PPTXParser().parse(presentation_path)

	assert "[Slide: 1]" in slides
	assert "KARYA PPTX test" in slides
