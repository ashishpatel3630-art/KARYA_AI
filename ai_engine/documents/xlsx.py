from pathlib import Path

from openpyxl import load_workbook


class XLSXParser:
    """Extract text and data from Excel XLSX documents."""

    def parse(self, file_path: str) -> str:
        """
        Extract data from all worksheets in an XLSX file.

        Args:
            file_path: Path to the XLSX file.

        Returns:
            Spreadsheet content as text.
        """

        path = Path(file_path)

        if not path.exists():
            raise FileNotFoundError(
                f"XLSX file not found: {file_path}"
            )

        if path.suffix.lower() != ".xlsx":
            raise ValueError(
                f"Expected an XLSX file, got: {path.suffix}"
            )

        workbook = load_workbook(
            filename=str(path),
            data_only=True,
        )

        content = []

        for worksheet in workbook.worksheets:
            content.append(f"[Sheet: {worksheet.title}]")

            for row in worksheet.iter_rows(values_only=True):
                values = []

                for value in row:
                    if value is not None:
                        values.append(str(value))

                if values:
                    content.append(" | ".join(values))

        return "\n".join(content)
